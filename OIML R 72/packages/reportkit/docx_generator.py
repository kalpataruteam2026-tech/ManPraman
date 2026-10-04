"""
WeighReport ReportKit - Word DOCX Generator
Builds fully editable Microsoft Word (.docx) reports from canonical ReportView.
Features genuine Word tables, styles, headers, footers, and failure formatting.
"""
import os
from typing import Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from .builder import get_test_title

def set_cell_shading(cell, color_hex: str):
    """Set cell background color."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def generate_docx_report(
    view: Dict[str, Any],
    output_path: str,
    kind: str = "detailed"
) -> str:
    """
    Renders ReportView into a Microsoft Word (.docx) document.
    """
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc = Document()

    # Page setup: Standard A4
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    # Header and Footer
    header = section.header
    hp = header.paragraphs[0]
    hp.text = "OIML R 76-2:2007 (E) Type Evaluation Report — Non-Automatic Weighing Instruments"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.style.font.size = Pt(8.5)
    hp.style.font.color.rgb = RGBColor(100, 116, 139)

    footer = section.footer
    fp = footer.paragraphs[0]
    meta = view.get("meta", {})
    fp.text = f"WeighReport System | App No: {meta.get('application_no')} | Rule Pack: {meta.get('rulepack_id')}"
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp.style.font.size = Pt(8.5)
    fp.style.font.color.rgb = RGBColor(100, 116, 139)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("OIML R 76-2 TYPE EVALUATION REPORT\n")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    rep_name = "SHORT SUMMARY EVALUATION REPORT" if kind == "summary" else "DETAILED TYPE EVALUATION REPORT"
    run_sub = p_title.add_run(f"{rep_name} — Non-Automatic Weighing Instruments")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Verdict Banner
    summary = view.get("summary", {})
    verdict = summary.get("overall_verdict", "INCOMPLETE")
    banner_color = "10B981" if verdict == "PASSED" else ("EF4444" if verdict == "FAILED" else "F59E0B")

    tbl_banner = doc.add_table(rows=1, cols=1)
    tbl_banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_b = tbl_banner.cell(0, 0)
    set_cell_shading(cell_b, banner_color)
    p_b = cell_b.paragraphs[0]
    p_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_b = p_b.add_run(f"OVERALL EVALUATION VERDICT: {verdict}")
    run_b.font.name = "Arial"
    run_b.font.size = Pt(13)
    run_b.font.bold = True
    run_b.font.color.rgb = RGBColor(255, 255, 255)

    doc.add_paragraph() # Spacer

    # Section 1: General Info
    p_h1 = doc.add_heading("1. General Information Concerning the Type", level=2)
    p_h1.style.font.color.rgb = RGBColor(30, 41, 59)

    inst = view.get("instrument", {})
    lab = view.get("lab", {})

    gen_rows = [
        ("Application No.:", str(meta.get("application_no")), "Date of Report:", str(meta.get("generated_at"))),
        ("Type Designation:", str(inst.get("type_designation")), "Accuracy Class:", f"Class {inst.get('accuracy_class')}"),
        ("Manufacturer:", str(inst.get("manufacturer")), "Max / Min:", f"{inst.get('Max')} {inst.get('unit')} / {inst.get('Min')} {inst.get('unit')}"),
        ("Applicant:", str(inst.get("applicant")), "e / d:", f"e = {inst.get('e')} {inst.get('unit')}, d = {inst.get('d')} {inst.get('unit')}"),
        ("Instrument Category:", str(inst.get("category")), "n = Max/e:", str(inst.get("n"))),
        ("Testing Lab:", str(lab.get("name", "National Metrology Lab")), "Accreditation:", str(lab.get("accreditation_no", "ISO 17025")))
    ]

    t_gen = doc.add_table(rows=len(gen_rows), cols=4)
    t_gen.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(gen_rows):
        for c_idx in range(4):
            cell = t_gen.cell(r_idx, c_idx)
            cell.text = row[c_idx]
            if c_idx % 2 == 0:
                cell.paragraphs[0].runs[0].font.bold = True
                set_cell_shading(cell, "F8FAFC")

    doc.add_paragraph() # Spacer

    # Section 2: Summary Table
    doc.add_heading("2. Summary of Type Evaluation (Pass/Fail Results)", level=2)
    sum_headers = ["Test No.", "Test Designation", "Verdict", "Key Metric / Margin", "Remarks"]
    sum_rows = summary.get("rows", [])

    t_sum = doc.add_table(rows=len(sum_rows) + 1, cols=5)
    t_sum.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Headers
    for c_idx, h in enumerate(sum_headers):
        cell = t_sum.cell(0, c_idx)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_shading(cell, "F1F5F9")

    # Data
    for r_idx, r in enumerate(sum_rows):
        row_cells = t_sum.rows[r_idx + 1].cells
        row_cells[0].text = str(r.get("test_no"))
        row_cells[1].text = str(r.get("test_name"))
        r_verdict = str(r.get("result"))
        row_cells[2].text = r_verdict
        run_v = row_cells[2].paragraphs[0].runs[0]
        run_v.font.bold = True
        if r_verdict == "PASSED":
            run_v.font.color.rgb = RGBColor(16, 185, 129)
        elif r_verdict == "FAILED":
            run_v.font.color.rgb = RGBColor(239, 68, 68)
            set_cell_shading(row_cells[2], "FEE2E2")
        row_cells[3].text = str(r.get("key_result"))
        row_cells[4].text = str(r.get("remarks") or "-")

    # Failures list
    failures = summary.get("failures", [])
    if failures:
        doc.add_paragraph()
        p_fl = doc.add_paragraph()
        run_fl = p_fl.add_run("Identified Non-Compliances / Failures:")
        run_fl.font.bold = True
        run_fl.font.color.rgb = RGBColor(220, 38, 38)
        for f in failures:
            doc.add_paragraph(f"• {f.get('test_code')} ({f.get('test_name')}): {f.get('reason')}")

    # Detailed section
    if kind == "detailed":
        doc.add_page_break()
        doc.add_heading("3. Test Observations and Detailed Calculations", level=2)

        for t in view.get("tests", []):
            code = t.get("test_code", "")
            title = get_test_title(code)
            t_verdict = str(t.get("verdict", "incomplete")).upper()

            p_th = doc.add_paragraph()
            r_th = p_th.add_run(f"{code} — {title}  [Verdict: {t_verdict}]")
            r_th.font.bold = True
            r_th.font.size = Pt(11)

            rows = t.get("results", {}).get("rows", []) or t.get("data", {}).get("rows", [])
            if rows and isinstance(rows, list):
                headers = ["Loc", "Kind", "Load L", "Indication I", "ΔL", "P", "E", "E0", "Ec", "mpe", "Result"]
                t_obs = doc.add_table(rows=len(rows) + 1, cols=len(headers))
                t_obs.alignment = WD_TABLE_ALIGNMENT.CENTER

                for c_idx, h in enumerate(headers):
                    cell = t_obs.cell(0, c_idx)
                    cell.text = h
                    cell.paragraphs[0].runs[0].font.bold = True
                    set_cell_shading(cell, "F8FAFC")

                for r_idx, r in enumerate(rows):
                    if not isinstance(r, dict):
                        continue
                    row_cells = t_obs.rows[r_idx + 1].cells
                    is_pass = r.get("pass", True)

                    row_cells[0].text = str(r.get("loc", r_idx + 1))
                    row_cells[1].text = str(r.get("kind", "load"))
                    row_cells[2].text = str(r.get("L", "-"))
                    row_cells[3].text = str(r.get("I", "-"))
                    row_cells[4].text = str(r.get("dL", "-"))
                    row_cells[5].text = str(r.get("P", "-"))
                    row_cells[6].text = str(r.get("E", "-"))
                    row_cells[7].text = str(r.get("E0", "-"))
                    row_cells[8].text = str(r.get("Ec", "-"))
                    row_cells[9].text = str(r.get("mpe", "-"))
                    row_cells[10].text = "PASS" if is_pass else "FAIL"

                    if not is_pass:
                        for cell in row_cells:
                            set_cell_shading(cell, "FEE2E2")
                        row_cells[10].paragraphs[0].runs[0].font.color.rgb = RGBColor(220, 38, 38)
                    else:
                        row_cells[10].paragraphs[0].runs[0].font.color.rgb = RGBColor(16, 185, 129)

            doc.add_paragraph() # Spacer

        # Methodology
        doc.add_page_break()
        doc.add_heading("4. Calculation Methodology and Rule Pack Specification", level=2)
        doc.add_paragraph(
            "This test report was evaluated automatically by the WeighReport Rule Engine using exact decimal rational arithmetic. "
            "All Maximum Permissible Error (mpe) limits are computed directly from OIML R 76-1:2006 Table 6 according to accuracy class, "
            "verification scale interval (e), and net test load. Operator verdict override is mathematically prohibited."
        )
        doc.add_paragraph(f"• Rule Pack: {meta.get('rulepack_id')} | Engine Version: {meta.get('engine_version')}")
        doc.add_paragraph(f"• Data Snapshot Hash: {meta.get('data_hash')}")

    # Signatures
    doc.add_paragraph()
    t_sign = doc.add_table(rows=3, cols=2)
    t_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sign.cell(0, 0).text = "Evaluated By (Test Engineer):"
    t_sign.cell(0, 0).paragraphs[0].runs[0].font.bold = True
    t_sign.cell(0, 1).text = "Approved By (Technical Reviewer):"
    t_sign.cell(0, 1).paragraphs[0].runs[0].font.bold = True

    t_sign.cell(1, 0).text = "Date: ________________________\nSignature: ___________________"
    t_sign.cell(1, 1).text = "Date: ________________________\nSignature: ___________________"
    t_sign.cell(2, 0).text = "Name: Metrology Engineer"
    t_sign.cell(2, 1).text = "Name: Metrology Technical Manager"

    doc.save(output_path)
    return output_path
