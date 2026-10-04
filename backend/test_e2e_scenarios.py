import urllib.request
import urllib.parse
import json
import io
import sys
import time

BASE_URL = "http://127.0.0.1:8000/api/v1"

def req(endpoint, method="GET", data=None, token=None, content_type="application/json"):
    url = f"{BASE_URL}{endpoint}"
    headers = {}
    body = None
    if token:
        headers["Authorization"] = f"Bearer {token}"
    if data is not None:
        if content_type == "application/json":
            headers["Content-Type"] = "application/json"
            body = json.dumps(data).encode("utf-8")
        else:
            headers["Content-Type"] = content_type
            body = data
    request = urllib.request.Request(url, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(request) as response:
            res_body = response.read()
            status_code = response.getcode()
            headers = dict(response.headers)
            try:
                return status_code, json.loads(res_body.decode("utf-8"))
            except Exception:
                return status_code, res_body
    except urllib.error.HTTPError as e:
        err_body = e.read().decode("utf-8")
        try:
            return e.code, json.loads(err_body)
        except Exception:
            return e.code, err_body

def run_tests():
    print("==================================================")
    print("RUNNING METROCERT E2E SUITE (SCENARIOS A through G)")
    print("==================================================")

    # 1. Check status
    code, status = req("/auth/status")
    print(f"[STATUS] Initial state: code={code}, status={status}")

    # TEST A: ADMIN BOOTSTRAP & USER MANAGEMENT
    print("\n--- TEST A: ADMIN BOOTSTRAP & USER PROVISIONING ---")
    admin_token = None
    if not status.get("initialized"):
        code, boot_res = req("/auth/bootstrap", method="POST", data={
            "full_name": "Chief Administrator",
            "email": "admin@metrocert.gov",
            "password": "AdminSecurePassword2026!"
        })
        print(f"Bootstrap Admin: code={code}, res={boot_res}")
        assert code == 200, f"Failed to bootstrap admin: {boot_res}"
        admin_token = boot_res["access_token"]
    else:
        code, login_res = req("/auth/login", method="POST", data={
            "email": "admin@metrocert.gov",
            "username": "admin@metrocert.gov",
            "password": "AdminSecurePassword2026!"
        })
        print(f"Login Admin: code={code}")
        assert code == 200
        admin_token = login_res["access_token"]

    assert admin_token is not None

    # Test invalid role creation (e.g., inspector, auditor)
    code, err_res = req("/users", method="POST", token=admin_token, data={
        "full_name": "Old Inspector",
        "email": "inspector@metrocert.gov",
        "password": "Password123!",
        "role": "inspector"
    })
    print(f"Attempting to create forbidden role 'inspector': code={code} (Expect 422)")
    assert code == 422, "System must reject roles other than admin, technician, reviewer!"

    # Create Technician
    code, tech_user = req("/users", method="POST", token=admin_token, data={
        "full_name": "Test Technician 1",
        "email": "tech1@metrocert.gov",
        "password": "TechSecurePassword2026!",
        "role": "technician"
    })
    print(f"Create Technician: code={code}, user={tech_user}")
    assert code in (200, 400) # 400 if already exists

    # Create Reviewer
    code, rev_user = req("/users", method="POST", token=admin_token, data={
        "full_name": "Senior Reviewer 1",
        "email": "reviewer1@metrocert.gov",
        "password": "ReviewerSecurePassword2026!",
        "role": "reviewer"
    })
    print(f"Create Reviewer: code={code}, user={rev_user}")
    assert code in (200, 400) # 400 if already exists

    # TEST B: TECHNICIAN LOGIN & TEST SESSION CREATION
    print("\n--- TEST B: TECHNICIAN SINGLE INPUT & EVALUATION ---")
    code, tech_login = req("/auth/login", method="POST", data={
        "email": "tech1@metrocert.gov",
        "username": "tech1@metrocert.gov",
        "password": "TechSecurePassword2026!"
    })
    assert code == 200, f"Technician login failed: {tech_login}"
    tech_token = tech_login["access_token"]
    print(f"Technician logged in successfully. User: {tech_login['user']}")
    assert tech_login['user']['role'] == "technician"

    # Register/Select Model
    app_no = f"APP-{int(time.time())}"
    code, app_record = req("/applications", method="POST", token=tech_token, data={
        "application_no": app_no,
        "type_designation": "PS-60K",
        "applicant": "Precise Scales Ltd",
        "manufacturer": "Precise Scales Ltd",
        "instrument_category": "Complete instrument"
    })
    print(f"Registered Model: code={code}, id={app_record.get('id')}")
    assert code == 200, f"Failed to register model: {app_record}"
    app_id = app_record["id"]

    # Run automated calculations on normalized observations
    calc_req = {
        "test_code": "T1",
        "instrument": {
            "accuracy_class": "III",
            "Max": 60.0,
            "e": 0.02,
            "d": 0.02,
            "unit": "kg"
        },
        "test_payload": {
            "header": {},
            "config": {},
            "data": {
                "rows": [
                    {"L": 0.0, "I": 0.0, "dL": 0.010},
                    {"L": 10.0, "I": 10.0, "dL": 0.010},
                    {"L": 20.0, "I": 20.0, "dL": 0.010},
                    {"L": 40.0, "I": 40.0, "dL": 0.010},
                    {"L": 60.0, "I": 60.0, "dL": 0.010}
                ]
            }
        }
    }
    code, calc_res = req("/tests/calc/preview", method="POST", token=tech_token, data=calc_req)
    print(f"Automated Calculations: code={code}, calc_res={calc_res}")
    assert code == 200, f"Calculation failed: {calc_res}"
    assert isinstance(calc_res, dict) and calc_res.get("verdict") in ("passed", "PASS")

    # Create report draft / session
    code, session = req("/reports", method="POST", token=tech_token, data={
        "application_id": app_id,
        "instrument": {
            "model": "PS-60K",
            "accuracy_class": "III",
            "max": 60.0,
            "min": 0.4,
            "e": 0.02,
            "d": 0.02,
            "unit": "kg"
        },
        "remarks": {
            "technician": "Test Technician 1",
            "notes": "Single-model UI test run under controlled laboratory conditions",
            "observations_raw": calc_req,
            "calculations_summary": calc_res
        }
    })
    print(f"Created Session draft: code={code}, id={session.get('id')}")
    assert code == 200, f"Failed to create report session: {session}"
    session_id = session["id"]

    # Submit for review
    code, submit_res = req(f"/reports/{session_id}/submit", method="POST", token=tech_token)
    print(f"Submitted Session for Review: code={code}, status={submit_res.get('report_status')}")
    assert code == 200, f"Submit failed: {submit_res}"
    assert submit_res["report_status"] == "under_review"

    # TEST C: REVIEWER VERIFICATION & RETURN FOR CORRECTION
    print("\n--- TEST C: REVIEWER INSPECTION & RETURN FOR CORRECTION ---")
    code, rev_login = req("/auth/login", method="POST", data={
        "email": "reviewer1@metrocert.gov",
        "username": "reviewer1@metrocert.gov",
        "password": "ReviewerSecurePassword2026!"
    })
    assert code == 200, f"Reviewer login failed: {rev_login}"
    rev_token = rev_login["access_token"]
    assert rev_login['user']['role'] == "reviewer"

    # Reviewer lists pending sessions
    code, pending = req("/reports?status=under_review", method="GET", token=rev_token)
    print(f"Pending review list: count={len(pending)}")
    assert any(s["id"] == session_id for s in pending)

    # Return for correction without reason -> Expect 400
    code, err_ret = req(f"/reports/{session_id}/return", method="POST", token=rev_token, data={
        "reason": "   "
    })
    print(f"Return without reason: code={code} (Expect 400)")
    assert code == 400

    # Return for correction with valid reason
    code, return_res = req(f"/reports/{session_id}/return", method="POST", token=rev_token, data={
        "reason": "Please attach environmental condition logs (temperature and humidity)."
    })
    print(f"Return for correction: code={code}, status={return_res.get('report_status')}")
    assert code == 200
    assert return_res["report_status"] == "returned"

    # TEST D: TECHNICIAN RESUBMISSION & REVIEWER APPROVAL
    print("\n--- TEST D: RESUBMISSION & FINAL APPROVAL ---")
    # Technician checks returned session and resubmits
    code, resubmit_res = req(f"/reports/{session_id}/submit", method="POST", token=tech_token)
    assert code == 200, f"Resubmit failed: {resubmit_res}"
    assert resubmit_res["report_status"] == "under_review"

    # Reviewer approves
    code, approve_res = req(f"/reports/{session_id}/approve", method="POST", token=rev_token)
    print(f"Approved by Reviewer: code={code}, status={approve_res.get('report_status')}")
    assert code == 200, f"Approval failed: {approve_res}"
    assert approve_res["report_status"] == "finalised"

    # TEST E: FINAL REPORT GENERATION & QR VERIFICATION
    print("\n--- TEST E: FINAL REPORT DOWNLOAD & QR VERIFICATION ---")
    code, pdf_bytes = req(f"/reports/{session_id}/pdf", method="GET", token=tech_token)
    print(f"Download PDF Report: code={code}, bytes_len={len(pdf_bytes)}")
    assert code == 200
    assert pdf_bytes.startswith(b"%PDF")

    code, verify_res = req(f"/verify/{session_id}", method="GET")
    print(f"QR Verification: code={code}, verified={verify_res.get('verified')}, status={verify_res.get('status')}")
    assert code == 200
    assert verify_res["verified"] is True
    assert verify_res["status"] in ("finalised", "approved")

    # TEST F: MULTI-MODEL WORKBOOK BATCH IMPORT
    print("\n--- TEST F: MULTI-MODEL WORKBOOK BATCH IMPORT ---")
    import openpyxl
    wb = openpyxl.Workbook()
    # Sheet 1: Model A (valid)
    ws1 = wb.active
    ws1.title = "Model_Alpha_15K"
    ws1.append(["Key", "Value"])
    ws1.append(["Model", "Alpha-15K"])
    ws1.append(["Manufacturer", "Alpha Metrology"])
    ws1.append(["Class", "III"])
    ws1.append(["Max", "15"])
    ws1.append(["Min", "0.1"])
    ws1.append(["e", "0.005"])
    ws1.append(["d", "0.005"])
    ws1.append(["Unit", "kg"])
    ws1.append(["Load", "Indication", "DeltaL"])
    ws1.append([0.0, 0.0, 0.002])
    ws1.append([5.0, 5.0, 0.002])
    ws1.append([10.0, 10.0, 0.002])
    ws1.append([15.0, 15.0, 0.002])

    # Sheet 2: Model B (valid)
    ws2 = wb.create_sheet(title="Model_Beta_30K")
    ws2.append(["Key", "Value"])
    ws2.append(["Model", "Beta-30K"])
    ws2.append(["Manufacturer", "Beta Instruments"])
    ws2.append(["Class", "III"])
    ws2.append(["Max", "30"])
    ws2.append(["Min", "0.2"])
    ws2.append(["e", "0.01"])
    ws2.append(["d", "0.01"])
    ws2.append(["Unit", "kg"])
    ws2.append(["Load", "Indication", "DeltaL"])
    ws2.append([0.0, 0.0, 0.005])
    ws2.append([10.0, 10.0, 0.005])
    ws2.append([20.0, 20.0, 0.005])
    ws2.append([30.0, 30.0, 0.005])

    # Sheet 3: Model C (invalid - missing max and observations)
    ws3 = wb.create_sheet(title="Model_Gamma_Invalid")
    ws3.append(["Key", "Value"])
    ws3.append(["Model", "Gamma-Bad"])

    wb_bytes = io.BytesIO()
    wb.save(wb_bytes)
    wb_bytes.seek(0)
    raw_xlsx = wb_bytes.getvalue()

    # Create multipart form-data
    boundary = "----WebKitFormBoundary7MA4YWxkTrZu0gW"
    body = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="file"; filename="Test_Batch.xlsx"\r\n'
        f"Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n"
    ).encode("utf-8") + raw_xlsx + f"\r\n--{boundary}--\r\n".encode("utf-8")

    # Preview
    code, preview_res = req("/import/preview", method="POST", token=tech_token, data=body, content_type=f"multipart/form-data; boundary={boundary}")
    sheets = preview_res.get("sheets", [])
    valid_count = sum(1 for s in sheets if s.get("status") in ("valid", "warning"))
    error_count = sum(1 for s in sheets if s.get("status") == "error")
    print(f"Batch Preview: code={code}, total_sheets={preview_res.get('total_sheets')}, valid={valid_count}, error={error_count}")
    assert code == 200
    assert preview_res["total_sheets"] == 3
    assert valid_count == 2
    assert error_count == 1

    # Process
    code, process_res = req("/import/process", method="POST", token=tech_token, data=body, content_type=f"multipart/form-data; boundary={boundary}")
    print(f"Batch Process: code={code}, success={process_res.get('success_count')}, failures={process_res.get('failure_count')}")
    assert code == 200
    assert process_res["success_count"] == 2
    assert process_res["failure_count"] == 1

    # TEST G: SECURITY & PERMISSION ENFORCEMENT
    print("\n--- TEST G: SECURITY & ROLE AUTHORIZATION GUARDS ---")
    # 1. Technician tries to access user management -> 403
    code, _ = req("/users", method="GET", token=tech_token)
    print(f"Technician -> GET /users: code={code} (Expect 403)")
    assert code == 403, "Technician must not access /users"

    # 2. Technician tries to approve own test session -> 403
    code, _ = req(f"/reports/{session_id}/approve", method="POST", token=tech_token)
    print(f"Technician -> POST /reports/{session_id}/approve: code={code} (Expect 403)")
    assert code == 403, "Technician must not approve test sessions"

    # 3. Reviewer tries to access user management -> 403
    code, _ = req("/users", method="GET", token=rev_token)
    print(f"Reviewer -> GET /users: code={code} (Expect 403)")
    assert code == 403, "Reviewer must not access /users"

    # 4. Reviewer tries to submit a test session -> 403
    code, _ = req(f"/reports/{session_id}/submit", method="POST", token=rev_token)
    print(f"Reviewer -> POST /reports/{session_id}/submit: code={code} (Expect 403)")
    assert code == 403, "Reviewer must not submit test sessions"

    print("\n==================================================")
    print("ALL E2E SCENARIOS (A through G) PASSED PERFECTLY!")
    print("==================================================")

if __name__ == "__main__":
    run_tests()
